"""Turn an uploaded packet into a normalised, source-checked packet.

    read (layout parser or Claude) -> validate shape (DRF serializer)
      -> ground every value against the document -> attach previews

Grounding is the guard against a model "reading" numbers that aren't there:
for each rate we look for its printed text on the page (or in the cell) it
claims to come from, record the box for highlighting, and mark it unverified
if it isn't found.
"""

from __future__ import annotations

import base64
import io
import os
import re
import threading

from .money import ExtractionError, money
from . import claude_extractor, layout_parser

PAGE_SCALE = 1.5  # preview image pixels per PDF point
# PDFium (used to render page previews) is not thread-safe, and the WSGI
# bridge serves requests on a thread pool, so renders are serialised.
_PDFIUM = threading.Lock()


def detect_format(filename: str, data: bytes) -> str:
    if data[:5] == b"%PDF-" or filename.lower().endswith(".pdf"):
        return "pdf"
    if data[:2] == b"PK" and filename.lower().endswith((".xlsx", ".xlsm")):
        return "xlsx"
    raise ExtractionError("Upload a PDF or an .xlsx spreadsheet")


def claude_available() -> bool:
    return bool(os.environ.get("ANTHROPIC_API_KEY"))


def read(data: bytes, fmt: str, method: str = "auto", client=None) -> tuple[dict, dict]:
    """Returns (raw packet, extractor info)."""
    if method == "auto":
        method = "claude" if claude_available() else "layout"
    if method == "claude":
        if client is None and not claude_available():
            raise ExtractionError("Reading with Claude needs ANTHROPIC_API_KEY on the server")
        raw, model = claude_extractor.extract(data, fmt, client=client)
        return raw, {"method": "claude", "model": model}
    parse = layout_parser.parse_pdf if fmt == "pdf" else layout_parser.parse_xlsx
    try:
        return parse(data), {"method": "layout", "model": None}
    except ExtractionError as e:
        hint = "" if claude_available() else " Set ANTHROPIC_API_KEY to read other carrier layouts with Claude."
        raise ExtractionError(f"{e}.{hint}") from e


def _norm_money(text: str) -> str | None:
    try:
        return str(money(text))
    except ExtractionError:
        return None


def ground(packet: dict, data: bytes, fmt: str) -> None:
    """Find every rate in the document; add bbox / verified flags in place."""
    if fmt == "pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            words_by_page = {i + 1: p.extract_words() for i, p in enumerate(pdf.pages)}
        for plan in packet["plans"]:
            for rate in plan["rates"]:
                src = rate.setdefault("source", {})
                want = rate["amount"]
                page = src.get("page")
                hits = [w for w in words_by_page.get(page, []) if _norm_money(w["text"]) == want] if page else []
                if src.get("bbox") and hits:
                    src["verified"] = True
                elif hits:
                    w = hits[0]
                    src.update(bbox=[round(w["x0"], 1), round(w["top"], 1), round(w["x1"], 1),
                                     round(w["bottom"], 1)], verified=True)
                else:
                    src["verified"] = False
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True)
        for plan in packet["plans"]:
            for rate in plan["rates"]:
                src = rate.setdefault("source", {})
                ok = False
                if src.get("sheet") in wb.sheetnames and src.get("cell"):
                    try:
                        cell = wb[src["sheet"]][src["cell"]]
                        ok = cell.value is not None and _norm_money(str(cell.value)) == rate["amount"]
                    except (ValueError, KeyError):
                        ok = False
                src["verified"] = ok


def pdf_pages(data: bytes) -> list[dict]:
    import pypdfium2 as pdfium

    pages = []
    with _PDFIUM:
        pdf = pdfium.PdfDocument(data)
        try:
            for i in range(len(pdf)):
                page = pdf[i]
                w, h = page.get_size()
                img = page.render(scale=PAGE_SCALE).to_pil()
                page.close()
                buf = io.BytesIO()
                img.convert("L").save(buf, format="PNG", optimize=True)
                pages.append({"n": i + 1, "width": round(w, 1), "height": round(h, 1),
                              "image": "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()})
        finally:
            pdf.close()
    return pages


def xlsx_sheets(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(max_row=80, max_col=12):
            rows.append([{"cell": c.coordinate, "value": None if c.value is None else str(c.value)}
                         for c in row])
        sheets.append({"name": ws.title, "rows": rows})
    return sheets


def first_age(label: str) -> int:
    from .money import band_ages

    try:
        return band_ages(label)[0]
    except ExtractionError:
        return 999


def plan_key(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
