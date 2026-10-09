"""Generate sample carrier renewal packets for the Renewal Build demo.

Run:  python scripts/make_renewal_packets.py
Writes data/renewal_packets/<group id>.(pdf|xlsx) plus <group id>.truth.json
(the values the packet was generated from, used by the tests).

Everything here is fictional: "Carrier A", the plan codes, the rates and the
people. Each packet is labelled as a sample on every page or sheet.

The renewal base rates match scripts/make_renewal_pdfs.py, so a packet and the
Renewal Advisor analysis of the same group agree to the cent. Each packet has
one deliberate problem for the build checks to catch:

  grp_30_2 (PDF)   Silver age 47 rate has two digits swapped (a keying error).
                   Silver HMO 3500 is renamed "Silver Select HMO 3500".
  grp_45_3 (XLSX)  Bronze HSA 6000 is discontinued and replaced by
                   "Bronze HSA 6000 Plus". Gold has no row for age 30.
  grp_12_1 (PDF)   Effective date printed as 02/01/2027 (Clasp expects 1/1),
                   and the packet's "current" Gold rate differs from Clasp's.
"""

from __future__ import annotations

import json
import sys
from datetime import date
from decimal import Decimal as D
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from renewal_advisor.age_curve import FEDERAL_DEFAULT  # noqa: E402
from renewal_advisor.census import generate_group  # noqa: E402
from renewal_advisor.rating import to_cents  # noqa: E402
from renewal_advisor.service import SAMPLES  # noqa: E402

OUT = ROOT / "data" / "renewal_packets"
SAMPLE_NOTE = "SAMPLE DOCUMENT - fictional carrier, plans and rates, generated for a demo."
INCREASES = {"plan_gold": D("1.19"), "plan_silver": D("1.18"), "plan_bronze": D("1.165")}
CODES = {"plan_gold": "CA-G1500", "plan_silver": "CA-S3500", "plan_bronze": "CA-B6000"}

# Age bands as carriers print them: one band for 0-14, single ages 15-63, 64+.
BANDS: list[tuple[str, int]] = [("0-14", 0)] + [(str(a), a) for a in range(15, 64)] + [("64+", 64)]


def band_rates(base21: D) -> list[tuple[str, D]]:
    return [(label, to_cents(base21 * FEDERAL_DEFAULT[age])) for label, age in BANDS]


def swap_two_digits(amount: D) -> D:
    """Simulate a keying error: swap the last two digits of the dollar part."""
    dollars, cents = f"{amount:.2f}".split(".")
    swapped = dollars[:-2] + dollars[-1] + dollars[-2]
    return D(f"{swapped}.{cents}")


def packet_spec(group_id: str) -> dict:
    g = generate_group(*{"grp_30_2": (30, 2), "grp_45_3": (45, 3), "grp_12_1": (12, 1)}[group_id])
    plans = []
    for p in g.plans:
        new21 = to_cents(p.base_rate_21 * INCREASES[p.id])
        plans.append({
            "current_plan_id": p.id, "current_plan_name": p.plan_name,
            "current_code": CODES[p.id] + "-26", "renewal_plan_name": p.plan_name,
            "renewal_code": CODES[p.id] + "-27", "status": "renews",
            "metal_level": p.metal_level, "deductible": str(p.deductible), "oop_max": str(p.oop_max),
            "plan_type": {"plan_gold": "ppo", "plan_silver": "hmo", "plan_bronze": "hdhp"}[p.id],
            "hsa_eligible": p.id == "plan_bronze",
            "current_rate_21": str(p.base_rate_21), "renewal_rate_21": str(new21),
            "rates": [[label, str(r)] for label, r in band_rates(new21)],
            "note": "Renews with no benefit changes.",
        })
    by_id = {p["current_plan_id"]: p for p in plans}
    effective = date(2027, 1, 1)
    problems: list[dict] = []

    if group_id == "grp_30_2":
        s = by_id["plan_silver"]
        s.update(renewal_plan_name="Silver Select HMO 3500", status="renamed",
                 note="Renamed for 2027 (Select network). Benefits unchanged.")
        i = next(i for i, (label, _) in enumerate(s["rates"]) if label == "47")
        correct = D(s["rates"][i][1])
        s["rates"][i][1] = str(swap_two_digits(correct))
        problems.append({"kind": "age_curve", "plan": "plan_silver", "age": "47",
                         "printed": s["rates"][i][1], "correct": str(correct)})
    elif group_id == "grp_45_3":
        b = by_id["plan_bronze"]
        b.update(renewal_plan_name="Bronze HSA 6000 Plus", status="replaced",
                 renewal_code="CA-B6000P-27",
                 note="Bronze HSA 6000 is discontinued. Members map to Bronze HSA 6000 Plus.")
        gold = by_id["plan_gold"]
        gold["rates"] = [r for r in gold["rates"] if r[0] != "30"]
        problems.append({"kind": "missing_age", "plan": "plan_gold", "age": "30"})
    elif group_id == "grp_12_1":
        effective = date(2027, 2, 1)
        clasp_rate = by_id["plan_gold"]["current_rate_21"]
        by_id["plan_gold"]["current_rate_21"] = "525.00"
        problems += [{"kind": "effective_date", "printed": "2027-02-01", "expected": "2027-01-01"},
                     {"kind": "current_rate", "plan": "plan_gold", "printed": "525.00",
                      "clasp": clasp_rate}]

    return {"group_id": group_id, "employer": SAMPLES[group_id].name, "state": g.state, "carrier": "Carrier A",
            "group_number": group_id.upper().replace("_", "-"),
            "effective_date": effective.isoformat(), "broker": "[Broker of record]",
            "reply_by": "2026-12-01", "plans": plans, "planted_problems": problems,
            "format": "xlsx" if group_id == "grp_45_3" else "pdf"}


def write_pdf(spec: dict, out: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    eff = date.fromisoformat(spec["effective_date"])
    head = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a5f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica-Oblique", 7.5)
        canvas.drawString(54, 30, SAMPLE_NOTE)
        canvas.drawRightString(558, 30, f"Page {doc.page}")
        canvas.restoreState()

    story = [
        Paragraph(f"{spec['carrier']} - Small Group Renewal", styles["Title"]),
        Paragraph(f"Employer: {spec['employer']}", styles["Normal"]),
        Paragraph(f"Group number: {spec['group_number']}", styles["Normal"]),
        Paragraph(f"Renewal effective: {eff:%m/%d/%Y}", styles["Normal"]),
        Paragraph(f"Rating state: {spec['state']}", styles["Normal"]),
        Paragraph(f"Please return plan selections by {date.fromisoformat(spec['reply_by']):%m/%d/%Y}.",
                  styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Plan changes for the new plan year", styles["Heading3"]),
    ]
    rows = [["Current plan", "Code", "Renewal plan", "Code", "Status", "Notes"]]
    for p in spec["plans"]:
        rows.append([p["current_plan_name"], p["current_code"], p["renewal_plan_name"],
                     p["renewal_code"], p["status"].capitalize(),
                     Paragraph(p["note"], styles["BodyText"])])
    t = Table(rows, colWidths=[82, 62, 100, 70, 52, 138])
    t.setStyle(head)
    story += [t, Spacer(1, 12), Paragraph("Rate summary (monthly, age 21)", styles["Heading3"])]
    rows = [["Renewal plan", "Deductible", "OOP max", "Current rate", "Renewal rate", "Change"]]
    for p in spec["plans"]:
        cur, new = D(p["current_rate_21"]), D(p["renewal_rate_21"])
        rows.append([p["renewal_plan_name"], f"${D(p['deductible']):,.0f}", f"${D(p['oop_max']):,.0f}",
                     f"${cur:,.2f}", f"${new:,.2f}", f"+{(new / cur - 1) * 100:.1f}%"])
    t = Table(rows, hAlign="LEFT")
    t.setStyle(head)
    story += [t, Spacer(1, 10), Paragraph(
        "Member premiums are age rated. Each covered person's rate is shown in the age-banded "
        "tables that follow; children: only the three oldest under 21 are charged.", styles["BodyText"])]

    for p in spec["plans"]:
        story += [PageBreak(), Paragraph(f"Age-banded monthly rates - {p['renewal_plan_name']}",
                                         styles["Heading2"]),
                  Paragraph(f"Plan code {p['renewal_code']}. Effective {eff:%m/%d/%Y}.",
                            styles["Normal"]), Spacer(1, 8)]
        rates = p["rates"]
        n = (len(rates) + 2) // 3
        cols = [rates[i * n:(i + 1) * n] for i in range(3)]
        rows = [["Age", "Monthly rate"] * 3]
        for r in range(n):
            row = []
            for c in cols:
                row += [c[r][0], f"${D(c[r][1]):,.2f}"] if r < len(c) else ["", ""]
            rows.append(row)
        t = Table(rows, colWidths=[44, 80] * 3, hAlign="LEFT")
        t.setStyle(head)
        story.append(t)

    out.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(out), pagesize=letter, title=f"{spec['carrier']} renewal - {spec['employer']}",
                      bottomMargin=54).build(story, onFirstPage=footer, onLaterPages=footer)


def write_xlsx(spec: dict, out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    eff = date.fromisoformat(spec["effective_date"])
    ws.append([f"{spec['carrier']} Small Group Renewal"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Employer", spec["employer"]])
    ws.append(["Group number", spec["group_number"]])
    ws.append(["Renewal effective", eff.strftime("%m/%d/%Y")])
    ws.append(["Rating state", spec["state"]])
    ws.append([])
    ws.append(["Current plan", "Current code", "Renewal plan", "Renewal code", "Status",
               "Deductible", "OOP max", "Current rate (21)", "Renewal rate (21)", "Notes"])
    for p in spec["plans"]:
        ws.append([p["current_plan_name"], p["current_code"], p["renewal_plan_name"], p["renewal_code"],
                   p["status"].capitalize(), float(p["deductible"]), float(p["oop_max"]),
                   float(p["current_rate_21"]), float(p["renewal_rate_21"]), p["note"]])
    ws.append([])
    ws.append([SAMPLE_NOTE])
    for p in spec["plans"]:
        sh = wb.create_sheet(p["renewal_plan_name"][:31])
        sh.append([f"Age-banded monthly rates - {p['renewal_plan_name']} ({p['renewal_code']})"])
        sh.append(["Age", "Monthly rate"])
        for label, amount in p["rates"]:
            sh.append([label, float(amount)])
        sh.append([])
        sh.append([SAMPLE_NOTE])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


if __name__ == "__main__":
    for gid in ["grp_30_2", "grp_45_3", "grp_12_1"]:
        spec = packet_spec(gid)
        path = OUT / f"{gid}.{spec['format']}"
        (write_xlsx if spec["format"] == "xlsx" else write_pdf)(spec, path)
        (OUT / f"{gid}.truth.json").write_text(json.dumps(spec, indent=2) + "\n")
        print("wrote", path.relative_to(ROOT))
